"""P1 — the setup pack's two directions (SETUP-REDESIGN-PLAN §4).

What each test is defending:

  * the blank pack a school receives carries NO example rows — a friendly
    template becomes a data-integrity bug the moment a school returns it with
    "Aarav Sharma" still in row 2;
  * every sheet round-trips: what the generator writes, the parser reads back
    identically, so the template can never drift from its reader (the rule
    `templates.py` set and `test_setup_onboarding.py` already asserts);
  * a REAL school's file still parses — the founder's reference sheet has a
    banner row above the headers, renames "Chapter" to "Chapter Name" and the
    teacher to "Mentor", and stores chapter numbers as floats;
  * "Class" and "Class teacher" sitting side by side do not collide, which is
    the whole reason this package has its own header mapper;
  * a blank cell stays None everywhere — an unsized chapter must never arrive
    as 1 and an unassigned subject must never arrive as 0;
  * a mid-year school (tracking start, half a syllabus) parses without
    complaint, because that is a state and not an error (D-3 / D-4).
"""

import io
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

from app.services.roster_import import parse_dob
from app.services.setup_pack import (
    COMMIT_ORDER,
    SHEETS,
    build_pack,
    filename_for,
    fill_down,
    parse_pack,
    sheet,
)

TODAY = date(2026, 8, 7)


def _book(sheets: dict[str, list[list]]) -> bytes:
    """A workbook built the way a school would hand it to us."""
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── the blank pack ───────────────────────────────────────────────────────────
def test_blank_pack_has_every_sheet_and_a_read_me():
    wb = load_workbook(io.BytesIO(build_pack("Sunrise Public School")))
    titles = wb.sheetnames
    assert titles[0] == "Read Me"
    for spec in SHEETS:
        assert spec.title in titles


def test_blank_pack_ships_no_example_rows():
    """The defence: a school that returns the pack untouched must import nothing,
    not three imaginary children."""
    wb = load_workbook(io.BytesIO(build_pack("Sunrise")))
    for spec in SHEETS:
        if spec.is_key_value:
            continue
        rows = list(wb[spec.title].iter_rows(min_row=2, values_only=True))
        assert not [r for r in rows if any(c is not None for c in r)], spec.title

    assert parse_pack(build_pack("Sunrise")).total_rows == 0


def test_blank_pack_prefills_what_the_operator_already_typed():
    pack = parse_pack(build_pack("Sunrise", overrides={
        "school_name": "Sunrise Public School", "state": "Karnataka"}))
    assert pack.setting("school_name") == "Sunrise Public School"
    assert pack.setting("state") == "Karnataka"
    # Nothing else invented on the school's behalf.
    assert pack.setting("year_label") is None


def test_filename_is_findable_in_a_folder_of_thirty():
    assert filename_for("Sunrise Public School") == (
        "TrackBit-Setup-Pack-Sunrise-Public-School.xlsx")
    assert filename_for(None).endswith(".xlsx")


# ── the round trip ───────────────────────────────────────────────────────────
def test_sample_pack_round_trips_every_sheet():
    pack = parse_pack(build_pack("Sunrise", sample=True))
    assert pack.missing_sheets == []
    for spec in SHEETS:
        data = pack.sheets[spec.key]
        assert data.present, spec.title
        assert data.missing_columns == [], spec.title
        if not spec.is_key_value:
            assert data.row_count == len(spec.samples), spec.title


def test_every_required_column_maps_from_its_own_header():
    """The anti-drift assertion: a template whose headers its own parser cannot
    place would send the school through the gap screen it exists to skip."""
    pack = parse_pack(build_pack("Sunrise", sample=True))
    for spec in SHEETS:
        data = pack.sheets[spec.key]
        for key in spec.required_keys:
            assert key in data.mapping, f"{spec.title}.{key}"


def test_school_settings_round_trip_including_the_mid_year_date():
    pack = parse_pack(build_pack("Sunrise", sample=True))
    assert pack.setting("year_label") == "2026-27"
    assert pack.setting("periods_per_day") == "8"
    # D-3: a school joining part-way through the year says so here.
    assert pack.setting("tracking_start") == "10/08/2026"


# ── what a real school's file looks like ─────────────────────────────────────
def test_banner_row_above_the_headers_is_skipped():
    """The reference file opens with 'CLASS 7 - SYLLABUS TRACKER (2026-27)'."""
    data = _book({"Syllabus": [
        ["CLASS 7 - SYLLABUS TRACKER (2026-27, Jul-Nov)"],
        ["Class", "Subject", "Term", "Ch #", "Chapter", "Periods"],
        ["7", "Mathematics", "Term 1", 1.0, "Integers", 4],
    ]})
    rows = parse_pack(data).rows("syllabus")
    assert len(rows) == 1
    assert rows[0]["chapter"] == "Integers"
    # A chapter numbered 1.0 reads as a bug to whoever reviews the import.
    assert rows[0]["chapter_no"] == "1"


def test_renamed_headers_still_map():
    """'Chapter Name' and 'Est. Periods' are the reference's own words."""
    data = _book({"Syllabus": [
        ["Class", "Subject", "Chapter Name", "Est. Periods"],
        ["7", "Science", "Nutrition in Plants", 4],
    ]})
    rows = parse_pack(data).rows("syllabus")
    assert rows[0]["chapter"] == "Nutrition in Plants"
    assert rows[0]["periods"] == "4"


def test_a_teacher_named_mentor_still_maps():
    data = _book({"Teaching Assignments": [
        ["Class", "Subject", "Mentor", "Periods per week"],
        ["7", "Maths", "Tejas Guru Ji", 6],
    ]})
    assert parse_pack(data).rows("assignments")[0]["teacher"] == "Tejas Guru Ji"


def test_class_and_class_teacher_do_not_collide():
    """The reason this package does not reuse ingest.heuristic_mapping: that one
    lets two fields claim the same column, and these two sit side by side."""
    data = _book({"Classes": [
        ["Class", "Section", "Class Teacher"],
        ["6", "A", "Anita Desai"],
    ]})
    row = parse_pack(data).rows("classes")[0]
    assert row["class_name"] == "6"
    assert row["class_teacher"] == "Anita Desai"


def test_a_blank_cell_is_none_and_never_a_number():
    """An unsized chapter is a state. Rendering it as 1 period would silently
    plan teaching nobody agreed to."""
    data = _book({"Syllabus": [
        ["Class", "Subject", "Chapter", "Periods"],
        ["6", "Mathematics", "Integers", None],
        ["6", "Mathematics", "Fractions", "  "],
    ]})
    rows = parse_pack(data).rows("syllabus")
    assert rows[0]["periods"] is None
    assert rows[1]["periods"] is None


def test_extra_columns_a_school_added_are_reported_not_dropped_silently():
    data = _book({"Syllabus": [
        ["Class", "Subject", "Chapter", "Difficulty", "CET Avg %"],
        ["7", "English", "A Hero", "Medium", 68.33],
    ]})
    syllabus = parse_pack(data).sheets["syllabus"]
    assert sorted(syllabus.unmapped_columns) == ["CET Avg %", "Difficulty"]


def test_date_cells_and_typed_dates_both_reach_parse_dob():
    data = _book({"Students": [
        ["Student name", "Admission no", "Date of birth"],
        ["Aarav Sharma", "1021", datetime(2014, 6, 14)],
        ["Diya Patel", "1022", "02/11/2013"],
    ]})
    rows = parse_pack(data).rows("students")
    assert rows[0]["date_of_birth"] == "2014-06-14"
    assert rows[1]["date_of_birth"] == "02/11/2013"
    # Both spellings survive into the parser that already owns this decision.
    assert parse_dob(rows[0]["date_of_birth"], today=TODAY) == date(2014, 6, 14)
    assert parse_dob(rows[1]["date_of_birth"], today=TODAY) == date(2013, 11, 2)


def test_sheet_titles_survive_case_and_dash_variants():
    data = _book({"TEACHING-ASSIGNMENTS": [
        ["Class", "Subject", "Teacher", "Periods per week"],
        ["6", "Mathematics", "Anita Desai", 6],
    ]})
    assert parse_pack(data).rows("assignments")[0]["teacher"] == "Anita Desai"


# ── what the school did not send ─────────────────────────────────────────────
def test_a_missing_required_sheet_is_reported_not_raised():
    pack = parse_pack(_book({"Students": [["Student name", "Admission no"],
                                          ["Aarav Sharma", "1021"]]}))
    assert "Syllabus" in pack.missing_sheets
    assert pack.sheets["syllabus"].present is False
    # An optional sheet's absence is not a missing sheet.
    assert "Fees" not in pack.missing_sheets


def test_a_missing_required_column_is_reported_per_sheet():
    data = _book({"Students": [["Student name", "Class"],
                               ["Aarav Sharma", "6"]]})
    assert parse_pack(data).sheets["students"].missing_columns == ["admission_no"]


def test_half_a_syllabus_parses_without_complaint():
    """D-4: a school that has only planned Terms 1 and 2 hands us those. The
    absent term is not an error — it simply is not there yet."""
    data = _book({"Syllabus": [
        ["Class", "Subject", "Term", "Chapter", "Periods"],
        ["6", "Mathematics", "Term 1", "Integers", 4],
        ["6", "Mathematics", "Term 2", "Algebra", None],
    ]})
    syllabus = parse_pack(data).sheets["syllabus"]
    assert syllabus.row_count == 2
    assert syllabus.missing_columns == []


def test_unknown_sheets_are_listed_and_the_read_me_is_not():
    data = _book({
        "Read Me": [["TrackBit setup pack"]],
        "Students": [["Student name", "Admission no"], ["Aarav Sharma", "1021"]],
        "Books & Coverage Flags": [["Class", "Note"]],
    })
    assert parse_pack(data).extra_sheets == ["Books & Coverage Flags"]


# ── merged cells ─────────────────────────────────────────────────────────────
def test_fill_down_carries_merged_cells_but_leaves_the_parsed_rows_alone():
    """One sheet now holds every class, so class and subject need the same
    blank-continues-above treatment chapter already had."""
    rows = [
        {"class_name": "6", "subject": "Mathematics", "chapter": "Integers",
         "topic": "Number line"},
        {"class_name": None, "subject": None, "chapter": None,
         "topic": "Absolute value"},
        {"class_name": None, "subject": None, "chapter": "Fractions",
         "topic": "Proper fractions"},
    ]
    merged = fill_down(rows, ("class_name", "subject", "chapter"))
    assert merged[1]["class_name"] == "6"
    assert merged[1]["chapter"] == "Integers"
    assert merged[2]["chapter"] == "Fractions"
    # The literal record of what the school sent is untouched.
    assert rows[1]["class_name"] is None


# ── the specs themselves ─────────────────────────────────────────────────────
def test_commit_order_covers_every_sheet_exactly_once():
    """The wizard's dependency chain moved into the order of writes; a sheet
    missing from it would simply never be committed."""
    assert sorted(COMMIT_ORDER) == sorted(s.key for s in SHEETS)


def test_the_syllabus_sheet_carries_class_and_subject():
    """D-5: one sheet for the whole school, not one file per class-subject."""
    keys = {c.key for c in sheet("syllabus").columns}
    assert {"class_name", "subject"} <= keys
    # Chapter and periods, nothing else (founder, 2026-08-08). Asking for a
    # topic got blank cells or the chapter name typed twice; the reference
    # tracker has no topic column at all.
    assert "topic" not in keys
    assert {"chapter", "periods"} <= keys
