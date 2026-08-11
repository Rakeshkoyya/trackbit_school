"""P2 — validating a filled pack (SETUP-REDESIGN-PLAN §7).

What each test is defending:

  * **the mid-year guarantee, in code.** An unsized chapter, an unplanned term
    and a subject with no teacher are NOTES. If any of them ever becomes a
    blocker, a school that has planned half its year can no longer be handed
    over — which is the exact promise D-3 and D-4 make;
  * **every finding names its consequence.** "45 parents cannot log in", not
    "45 nulls" — the operator is deciding whether to send this back to a
    principal;
  * **cross-sheet references really are checked**, because one sheet now holds
    every class: a syllabus row for a class nobody teaches would otherwise
    import into nothing;
  * **a blank Section fans out** to every section of the class, which is what
    makes one syllabus sheet enough (D-5);
  * **nothing is silently truncated.** A rule that fires 200 times reports the
    first few and then says how many more there are.
"""

import io

from openpyxl import Workbook

from app.services.setup_pack import (
    BLOCKER,
    NOTE,
    WARNING,
    ParsedPack,
    build_pack,
    parse_pack,
    validate,
)

SCHOOL_ROWS = [
    ["Field", "Value"],
    ["School name", "Sunrise Public School"],
    ["Academic year", "2026-27"],
    ["Year starts", "01/06/2026"],
    ["Year ends", "30/04/2027"],
    ["Periods per day", 8],
    ["Working days", "Mon-Sat"],
]
TERMS_ROWS = [
    ["Term", "Starts", "Ends"],
    ["Term 1", "01/06/2026", "30/09/2026"],
    ["Term 2", "01/10/2026", "30/04/2027"],
]
CLASSES_ROWS = [
    ["Class", "Section", "Class teacher"],
    ["6", "A", "Anita Desai"],
]
STAFF_ROWS = [
    ["Name", "Email", "Role"],
    ["Anita Desai", "anita@school.example", "Teacher"],
]
ASSIGNMENT_ROWS = [
    ["Class", "Section", "Subject", "Teacher"],
    ["6", "A", "Mathematics", "Anita Desai"],
]
SYLLABUS_ROWS = [
    ["Class", "Section", "Subject", "Term", "Chapter", "Periods"],
    ["6", "A", "Mathematics", "Term 1", "Integers", 4],
]
STUDENT_ROWS = [
    ["Student name", "Admission no", "Class", "Section", "Date of birth"],
    ["Aarav Sharma", "1021", "6", "A", "14/06/2014"],
]

TITLES = {
    "school": "School", "terms": "Terms", "classes": "Classes",
    "staff": "Staff", "assignments": "Teaching Assignments",
    "syllabus": "Syllabus", "students": "Students", "timetable": "Timetable",
    "calendar": "Calendar", "fees": "Fees",
}


def _pack(**overrides: list[list]) -> ParsedPack:
    """A minimally valid pack, with named sheets swapped out per test."""
    sheets = {
        "School": SCHOOL_ROWS, "Terms": TERMS_ROWS, "Classes": CLASSES_ROWS,
        "Staff": STAFF_ROWS, "Teaching Assignments": ASSIGNMENT_ROWS,
        "Syllabus": SYLLABUS_ROWS, "Students": STUDENT_ROWS,
    }
    for key, rows in overrides.items():
        sheets[TITLES[key]] = rows
    wb = Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    return parse_pack(buf.getvalue())


def _rules(report, severity: str | None = None) -> set[str]:
    return {f.rule for f in report.findings
            if severity is None or f.severity == severity}


# ── the happy path ───────────────────────────────────────────────────────────
def test_a_minimal_valid_pack_has_no_blockers():
    report = validate(_pack())
    assert report.ready, [f.message for f in report.blockers]


def test_the_filled_example_pack_has_no_blockers():
    """The pack an operator shows a school must not model a mistake."""
    report = validate(parse_pack(build_pack("Sunrise", sample=True)))
    assert report.ready, [f.message for f in report.blockers]


# ── a column we could not place (silent data loss until it is reported) ──────
def test_an_unrecognised_column_is_named_rather_than_dropped_in_silence():
    """The commonest real failure: the school heads its guardian columns its own
    way, every hint misses, and the review comes back green while every parent
    contact is discarded. `unmapped_columns` was computed and read by nobody."""
    rows = [
        ["Student name", "Admission no", "Class", "Section", "Date of birth",
         "Guardian Name", "Parent Contact"],
        ["Aarav Sharma", "1021", "6", "A", "14/06/2014", "Rajesh Sharma", "9876543210"],
    ]
    report = validate(_pack(students=rows))

    assert report.ready, "an unknown column is worth a warning, never a blocker"
    named = " ".join(f.message for f in report.findings
                     if f.severity == WARNING and f.rule == "students:unmapped_column")
    assert "Guardian Name" in named
    assert "Parent Contact" in named


def test_an_unrecognised_school_setting_is_reported_only_when_it_carries_a_value():
    """Column A of the School sheet also holds spacers and headings. A label with
    nothing beside it is one of those; a label with a value is a fact we dropped."""
    school = [*SCHOOL_ROWS, ["Principal", "Mrs Rao"], ["Notes", None]]
    messages = [f.message for f in validate(_pack(school=school)).findings
                if f.rule == "school:unmapped_column"]

    assert any("Principal" in m for m in messages)
    assert not any("Notes" in m for m in messages)


# ── month names (one date vocabulary, shared with the calendar importer) ──────
def test_month_name_dates_are_read_on_every_sheet():
    """A school that writes `30-Apr-2027` used to have the Calendar sheet accept
    it and the Terms sheet refuse it, because three importers owned three format
    lists. They share `core.dates` now."""
    school = [
        ["Field", "Value"],
        ["School name", "Sunrise Public School"],
        ["Academic year", "2026-27"],
        ["Year starts", "1 June 2026"],
        ["Year ends", "30-Apr-2027"],
        ["Periods per day", 8],
        ["Working days", "Mon-Sat"],
    ]
    terms = [
        ["Term", "Starts", "Ends"],
        ["Term 1", "01/06/2026", "30 Sep 2026"],
        ["Term 2", "01/10/2026", "30th April 2027"],
    ]
    report = validate(_pack(school=school, terms=terms))
    assert report.ready, [f.message for f in report.blockers]


def test_us_ordered_dates_are_still_refused():
    """Widening the vocabulary must not widen it to `%m/%d/%Y`: on the rows where
    both readings parse it turns 3 April into 4 March, silently."""
    school = [r for r in SCHOOL_ROWS if r[0] != "Year starts"]
    school.append(["Year starts", "06/21/2026"])
    assert "school:year_start" in _rules(validate(_pack(school=school)), BLOCKER)


# ── the mid-year guarantee (D-3 / D-4) ───────────────────────────────────────
def test_unsized_chapters_are_a_note_never_a_blocker():
    report = validate(_pack(syllabus=[
        ["Class", "Section", "Subject", "Term", "Chapter", "Periods"],
        ["6", "A", "Mathematics", "Term 1", "Integers", 4],
        ["6", "A", "Mathematics", "Term 1", "Fractions", None],
    ]))
    assert report.ready
    assert "syllabus:unsized" in _rules(report, NOTE)


def test_an_unplanned_term_is_a_note_never_a_blocker():
    """The school planned Term 1 and will send Term 2 in December."""
    report = validate(_pack())
    assert report.ready
    assert "syllabus:term_unplanned" in _rules(report, NOTE)
    note = next(f for f in report.findings if f.rule == "syllabus:term_unplanned")
    assert "Term 2" in note.message


def test_a_subject_with_no_teacher_is_a_note_never_a_blocker():
    report = validate(_pack(assignments=[
        ["Class", "Section", "Subject", "Teacher"],
        ["6", "A", "Mathematics", ""],
    ]))
    assert report.ready
    assert "assignments:no_teacher" in _rules(report, NOTE)


def test_a_subject_left_off_the_timetable_is_a_note():
    """TT-3's replacement for "no periods per week".

    That column is gone, so the state it described — *this subject has no
    weekly load, its plan cannot be paced* — is now visible as a subject the
    grid never mentions. Still a NOTE, never a blocker: a school halfway
    through drawing its timetable must still be handed over (D-3/D-4).
    """
    report = validate(_pack(
        assignments=[
            ["Class", "Section", "Subject", "Teacher"],
            ["6", "A", "Mathematics", "Anita Desai"],
            ["6", "A", "Science", "Anita Desai"]],
        syllabus=[["Class", "Section", "Subject", "Term", "Chapter", "Periods"],
                  ["6", "A", "Mathematics", "Term 1", "Integers", 4],
                  ["6", "A", "Science", "Term 1", "Food", 3]],
        timetable=[["Class", "Section", "Day", "Period", "Subject"],
                   ["6", "A", "Monday", 1, "Mathematics"]],
    ))
    assert report.ready
    note = next(f for f in report.findings
                if f.rule == "timetable:subject_not_scheduled")
    assert "Science" in note.message
    # Maths IS on the grid, so it is not reported.
    assert not any(f.rule == "timetable:subject_not_scheduled"
                   and "Mathematics" in f.message for f in report.findings)


def test_tracking_outside_the_year_warns_but_does_not_block():
    report = validate(_pack(school=SCHOOL_ROWS + [
        ["Tracking starts from", "10/08/2029"]]))
    assert report.ready
    assert "school:tracking_range" in _rules(report, WARNING)


# ── cross-sheet references ───────────────────────────────────────────────────
def test_a_syllabus_row_for_a_subject_nobody_teaches_is_blocked():
    report = validate(_pack(syllabus=[
        ["Class", "Section", "Subject", "Chapter"],
        ["6", "A", "Sanskrit", "Chapter 1"],
    ]))
    assert not report.ready
    assert "syllabus:not_offered" in _rules(report, BLOCKER)


def test_a_syllabus_row_for_a_class_that_does_not_exist_is_blocked():
    report = validate(_pack(syllabus=[
        ["Class", "Section", "Subject", "Chapter"],
        ["9", "C", "Mathematics", "Chapter 1"],
    ]))
    assert "syllabus:unknown_class" in _rules(report, BLOCKER)


def test_a_class_teacher_who_is_not_on_the_staff_sheet_is_blocked():
    report = validate(_pack(classes=[
        ["Class", "Section", "Class teacher"],
        ["6", "A", "Someone Else"],
    ]))
    blocker = next(f for f in report.blockers
                   if f.rule == "classes:unknown_teacher")
    # The consequence, not the field.
    assert "absence follow-ups" in blocker.message


def test_an_assignment_naming_an_unknown_teacher_is_blocked():
    report = validate(_pack(assignments=[
        ["Class", "Section", "Subject", "Teacher"],
        ["6", "A", "Mathematics", "Ghost Teacher", 6],
    ]))
    assert "assignments:unknown_teacher" in _rules(report, BLOCKER)


def test_a_class_with_no_subjects_is_flagged():
    report = validate(_pack(classes=[
        ["Class", "Section", "Class teacher"],
        ["6", "A", "Anita Desai"],
        ["7", "A", "Anita Desai"],
    ]))
    assert "assignments:class_without_subjects" in _rules(report, WARNING)


# ── a blank section means every section (D-5) ────────────────────────────────
def test_a_blank_section_fans_out_to_every_section_of_the_class():
    """One syllabus row for class 6 must satisfy 6-A and 6-B alike — that is
    what makes a single sheet enough for the whole school."""
    report = validate(_pack(
        classes=[["Class", "Section", "Class teacher"],
                 ["6", "A", "Anita Desai"],
                 ["6", "B", "Anita Desai"]],
        assignments=[
            ["Class", "Section", "Subject", "Teacher"],
            ["6", "", "Mathematics", "Anita Desai"]],
        syllabus=[["Class", "Section", "Subject", "Term", "Chapter", "Periods"],
                  ["6", "", "Mathematics", "Term 1", "Integers", 4]],
    ))
    assert report.ready, [f.message for f in report.blockers]
    assert "assignments:class_without_subjects" not in _rules(report)


# ── duplicates and unreadable values ─────────────────────────────────────────
def test_two_children_sharing_an_admission_number_is_blocked():
    report = validate(_pack(students=[
        ["Student name", "Admission no", "Class", "Section", "Date of birth"],
        ["Aarav Sharma", "1021", "6", "A", "14/06/2014"],
        ["Diya Patel", "1021", "6", "A", "02/11/2013"],
    ]))
    blocker = next(f for f in report.blockers
                   if f.rule == "students:duplicate_admission")
    assert "re-upload could not tell them apart" in blocker.message


def test_a_missing_date_of_birth_warns_in_parents_not_in_nulls():
    report = validate(_pack(students=[
        ["Student name", "Admission no", "Class", "Section", "Date of birth"],
        ["Aarav Sharma", "1021", "6", "A", ""],
    ]))
    warning = next(f for f in report.findings if f.rule == "students:dob")
    assert "cannot log in" in warning.message
    assert report.ready  # a warning, not a blocker


def test_a_missing_required_sheet_blocks():
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Students")
    for row in STUDENT_ROWS:
        ws.append(row)
    buf = io.BytesIO()
    wb.save(buf)
    report = validate(parse_pack(buf.getvalue()))
    assert not report.ready
    assert "syllabus:missing_sheet" in _rules(report, BLOCKER)


# ── the new_org invariants ───────────────────────────────────────────────────
# TT-3 moved both of these off Teaching Assignments and onto the Timetable
# sheet. They used to sum a "Periods per week" column a human typed; that column
# is gone, because a school filling the pack in April does not know what June's
# grid will give a subject. Counting the grid is also the stronger check: it
# tests the timetable the school will actually run, not its intention for one.
#
# The pack's week holds 8 periods/day × 6 days = 48.
_DAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday", "Saturday",
         "Sunday"]


def test_a_class_with_more_periods_on_the_grid_than_the_week_holds_warns():
    report = validate(_pack(timetable=[
        ["Class", "Section", "Day", "Period", "Subject"],
        # 7 days × 8 periods = 56, over the 48 the week holds.
        *[["6", "A", day, period, "Mathematics"]
          for day in _DAYS for period in range(1, 9)],
    ]))
    warning = next(f for f in report.findings
                   if f.rule == "timetable:class_over_capacity")
    assert "cannot run" in warning.message


def test_a_teacher_on_the_grid_more_than_the_week_holds_warns():
    """Counted across every class she teaches — the case the old column could
    only catch by adding up numbers nobody had checked against a real grid."""
    report = validate(_pack(
        classes=[["Class", "Section", "Class teacher"],
                 ["6", "A", "Anita Desai"],
                 ["7", "A", "Anita Desai"]],
        assignments=[
            ["Class", "Section", "Subject", "Teacher"],
            ["6", "A", "Mathematics", "Anita Desai"],
            ["7", "A", "Mathematics", "Anita Desai"]],
        syllabus=[["Class", "Section", "Subject", "Term", "Chapter", "Periods"],
                  ["6", "A", "Mathematics", "Term 1", "Integers", 4],
                  ["7", "A", "Mathematics", "Term 1", "Integers", 4]],
        timetable=[
            ["Class", "Section", "Day", "Period", "Subject"],
            # 28 periods in each of two classes — 56 for her, over the 48 that
            # exist in anybody's week.
            *[[c, "A", day, period, "Mathematics"]
              for c in ("6", "7")
              for day in _DAYS[:4] for period in range(1, 8)],
        ],
    ))
    assert "timetable:teacher_over_capacity" in _rules(report, WARNING)


# ── optional sheets ──────────────────────────────────────────────────────────
def test_a_double_booked_period_is_blocked():
    report = validate(_pack(timetable=[
        ["Class", "Section", "Day", "Period", "Subject"],
        ["6", "A", "Monday", 1, "Mathematics"],
        ["6", "A", "Monday", 1, "Mathematics"],
    ]))
    assert "timetable:double_booked" in _rules(report, BLOCKER)


def test_a_period_beyond_the_school_day_is_blocked():
    report = validate(_pack(timetable=[
        ["Class", "Section", "Day", "Period", "Subject"],
        ["6", "A", "Monday", 12, "Mathematics"],
    ]))
    assert "timetable:period_range" in _rules(report, BLOCKER)


def test_installments_that_do_not_sum_to_the_total_warn():
    report = validate(_pack(fees=[
        ["Class", "Category", "Total amount", "Installment no", "Due date",
         "Amount"],
        ["6", "", 48000, 1, "15/06/2026", 12000],
    ]))
    warning = next(f for f in report.findings if f.rule == "fees:installments_sum")
    assert "never collected" in warning.message


def test_an_absent_optional_sheet_is_not_a_finding():
    report = validate(_pack())
    assert "fees:missing_sheet" not in _rules(report)
    assert "timetable:missing_sheet" not in _rules(report)


# ── nothing is silently truncated ────────────────────────────────────────────
def test_a_rule_firing_many_times_says_how_many_it_did_not_list():
    rows = [["Student name", "Admission no", "Class", "Section", "Date of birth"]]
    rows += [[f"Child {i}", str(2000 + i), "9", "Z", "14/06/2014"]
             for i in range(20)]
    report = validate(_pack(students=rows))
    overflow = [f for f in report.findings
                if f.rule == "students:unknown_class:overflow"]
    assert overflow, "a capped rule must say how many more there were"
    assert "12 more" in overflow[0].message


def test_findings_are_ordered_worst_first():
    report = validate(_pack(syllabus=[
        ["Class", "Section", "Subject", "Chapter"],
        ["9", "C", "Mathematics", "Chapter 1"],
    ]))
    severities = [f.severity for f in report.findings]
    assert severities == sorted(
        severities, key=lambda s: {BLOCKER: 0, WARNING: 1, NOTE: 2}[s])


def test_a_blocking_row_reports_the_excel_row_the_school_must_open():
    report = validate(_pack(syllabus=[
        ["Class", "Section", "Subject", "Chapter"],
        ["6", "A", "Mathematics", "Integers"],
        ["9", "C", "Mathematics", "Chapter 1"],
    ]))
    blocker = next(f for f in report.blockers if f.rule == "syllabus:unknown_class")
    assert blocker.row == 3
