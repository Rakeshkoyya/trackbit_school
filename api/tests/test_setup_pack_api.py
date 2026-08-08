"""P5 — the operator's setup screen, on the wire (SETUP-REDESIGN-PLAN §5).

What each test is defending:

  * **only the operator can do any of this.** Schools stopped self-onboarding by
    founder decision on 2026-07-20, but until now that lived in the navigation
    only — a school admin who knew the URL could still run setup. These three
    routes are super-admin on the wire;
  * **review writes nothing.** The operator uploads, reads the report, sends the
    sheet back to the school, uploads again. If reviewing a broken pack left
    half a school behind, the whole loop would be unusable;
  * **a blocker refuses the import** and returns the report instead of building
    a school wrong;
  * **importing twice does not double the school** — the operator will do this;
  * **staff credentials come back once**, because the operator has to hand them
    over and they are hashed on the way in.
"""

import io
import uuid

from openpyxl import Workbook, load_workbook

from app.models import User
from app.services.setup_pack import build_pack
from tests.conftest import AdminSession

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


def _register(client, cleanup, org_name="Pack API School"):
    email = f"op-{uuid.uuid4().hex[:12]}@example.com"
    reg = client.post("/api/v1/auth/register-org", json={
        "org_name": org_name, "name": "Operator", "email": email,
        "password": "supersecret1", "timezone": "Asia/Kolkata"}).json()
    cleanup["orgs"].append(uuid.UUID(reg["org"]["id"]))
    cleanup["users"].append(uuid.UUID(reg["user"]["id"]))
    return reg


def _make_super(user_id: str) -> None:
    db = AdminSession()
    try:
        db.get(User, uuid.UUID(user_id)).is_super_admin = True
        db.commit()
    finally:
        db.close()


def _headers(reg):
    return {"Authorization": f"Bearer {reg['access_token']}"}


def _operator(client, cleanup):
    """A super-admin, and the id of a school they can set up."""
    reg = _register(client, cleanup)
    _make_super(reg["user"]["id"])
    return reg, reg["org"]["id"]


def _upload(data: bytes, name: str = "pack.xlsx"):
    return {"file": (name, io.BytesIO(data), XLSX)}


def _broken_pack() -> bytes:
    """Students only — every other required sheet missing."""
    wb = Workbook()
    wb.remove(wb.active)
    ws = wb.create_sheet("Students")
    ws.append(["Student name", "Admission no"])
    ws.append(["Aarav Sharma", "1021"])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ── the template ─────────────────────────────────────────────────────────────
def test_the_operator_downloads_a_pack_named_for_the_school(client, cleanup):
    reg, org_id = _operator(client, cleanup)
    r = client.get(f"/api/v1/platform/orgs/{org_id}/setup/template",
                   headers=_headers(reg))
    assert r.status_code == 200
    assert "Pack-API-School" in r.headers["content-disposition"]

    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames[0] == "Read Me"
    assert "Syllabus" in wb.sheetnames
    # Pre-filled with what the operator already typed, so nobody is asked twice.
    values = {row[0].value: row[1].value
              for row in wb["School"].iter_rows(min_row=2, max_col=2)}
    assert values["School name"] == "Pack API School"


def test_the_downloaded_pack_has_no_rows_to_delete(client, cleanup):
    reg, org_id = _operator(client, cleanup)
    r = client.get(f"/api/v1/platform/orgs/{org_id}/setup/template",
                   headers=_headers(reg))
    wb = load_workbook(io.BytesIO(r.content))
    rows = list(wb["Students"].iter_rows(min_row=2, values_only=True))
    assert not [row for row in rows if any(c is not None for c in row)]


# ── review writes nothing ────────────────────────────────────────────────────
def test_reviewing_a_good_pack_reports_ready_and_writes_nothing(client, cleanup):
    reg, org_id = _operator(client, cleanup)
    h = _headers(reg)
    r = client.post(f"/api/v1/platform/orgs/{org_id}/setup/review",
                    files=_upload(build_pack("Sunrise", sample=True)), headers=h)
    assert r.status_code == 200
    body = r.json()
    assert body["ready"] is True
    assert body["blockers"] == 0
    assert body["total_rows"] > 0

    # Nothing was written: the school still has no classes.
    assert client.get("/api/v1/academics/classes", headers=h).json() == []


def test_reviewing_a_broken_pack_names_the_missing_sheets(client, cleanup):
    reg, org_id = _operator(client, cleanup)
    r = client.post(f"/api/v1/platform/orgs/{org_id}/setup/review",
                    files=_upload(_broken_pack()), headers=_headers(reg))
    body = r.json()
    assert body["ready"] is False
    assert body["blockers"] > 0
    assert "Syllabus" in body["missing_sheets"]
    # Every finding carries something to say about it.
    assert all(f["message"] for f in body["findings"])


def test_a_file_that_is_not_a_workbook_is_refused_kindly(client, cleanup):
    reg, org_id = _operator(client, cleanup)
    r = client.post(f"/api/v1/platform/orgs/{org_id}/setup/review",
                    files={"file": ("notes.pdf", io.BytesIO(b"%PDF-1.4 nope"),
                                    "application/pdf")},
                    headers=_headers(reg))
    # 422 is what every AppError ValidationError returns here — not 400.
    assert r.status_code == 422
    assert "setup pack" in r.json()["error"]["message"]


# ── import ───────────────────────────────────────────────────────────────────
def test_importing_a_good_pack_builds_the_school(client, cleanup):
    reg, org_id = _operator(client, cleanup)
    h = _headers(reg)
    r = client.post(f"/api/v1/platform/orgs/{org_id}/setup/import",
                    files=_upload(build_pack("Sunrise", sample=True)), headers=h)
    assert r.status_code == 200
    body = r.json()
    assert body["imported"] is True

    created = {s["key"]: s["created"] for s in body["sheets"]}
    assert created["classes"] == 3
    assert created["students"] == 3
    assert created["terms"] == 2  # a term per exam window
    assert created["assignments"] >= 3

    # And the school is really there, read back through the ordinary API.
    assert len(client.get("/api/v1/academics/classes", headers=h).json()) == 3


def test_the_import_hands_back_the_staff_logins_once(client, cleanup):
    reg, org_id = _operator(client, cleanup)
    body = client.post(f"/api/v1/platform/orgs/{org_id}/setup/import",
                       files=_upload(build_pack("Sunrise", sample=True)),
                       headers=_headers(reg)).json()
    names = {c["name"] for c in body["credentials"]}
    assert {"Anita Desai", "Vikram Rao", "Sunita Iyer"} <= names
    assert all(c["username"] and c["password"] for c in body["credentials"])


def test_a_blocker_refuses_the_import_and_returns_the_report(client, cleanup):
    reg, org_id = _operator(client, cleanup)
    h = _headers(reg)
    r = client.post(f"/api/v1/platform/orgs/{org_id}/setup/import",
                    files=_upload(_broken_pack()), headers=h)
    assert r.status_code == 200
    body = r.json()
    assert body["imported"] is False
    assert body["review"]["blockers"] > 0
    assert body["sheets"] == []
    # A refused import builds nothing.
    assert client.get("/api/v1/academics/classes", headers=h).json() == []


def test_importing_twice_does_not_double_the_school(client, cleanup):
    reg, org_id = _operator(client, cleanup)
    h = _headers(reg)
    pack = build_pack("Sunrise", sample=True)
    for _ in range(2):
        assert client.post(f"/api/v1/platform/orgs/{org_id}/setup/import",
                           files=_upload(pack), headers=h).json()["imported"]
    assert len(client.get("/api/v1/academics/classes", headers=h).json()) == 3


def test_readiness_reads_the_school_the_pack_built(client, cleanup):
    reg, org_id = _operator(client, cleanup)
    h = _headers(reg)
    client.post(f"/api/v1/platform/orgs/{org_id}/setup/import",
                files=_upload(build_pack("Sunrise", sample=True)), headers=h)
    report = client.get(f"/api/v1/platform/orgs/{org_id}/readiness",
                        headers=h).json()
    checks = {c["key"]: c for c in report["checks"]}
    assert checks["classes"]["status"] == "ok"
    assert checks["students"]["status"] == "ok"
    # And handover can then be stamped.
    handed = client.post(f"/api/v1/platform/orgs/{org_id}/handover",
                         headers=h).json()
    assert handed["handed_over_at"] is not None


# ── only the operator ────────────────────────────────────────────────────────
def test_a_school_admin_cannot_reach_any_of_the_setup_routes(client, cleanup):
    """Schools do not self-onboard. Until now that was navigation only."""
    reg = _register(client, cleanup)  # an ordinary admin, NOT super
    org_id = reg["org"]["id"]
    h = _headers(reg)

    assert client.get(f"/api/v1/platform/orgs/{org_id}/setup/template",
                      headers=h).status_code == 403
    assert client.post(f"/api/v1/platform/orgs/{org_id}/setup/review",
                       files=_upload(build_pack("x", sample=True)),
                       headers=h).status_code == 403
    assert client.post(f"/api/v1/platform/orgs/{org_id}/setup/import",
                       files=_upload(build_pack("x", sample=True)),
                       headers=h).status_code == 403


def test_an_unknown_school_is_a_clean_404(client, cleanup):
    reg, _org_id = _operator(client, cleanup)
    r = client.get(f"/api/v1/platform/orgs/{uuid.uuid4()}/setup/template",
                   headers=_headers(reg))
    assert r.status_code == 404
