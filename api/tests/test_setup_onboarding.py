"""V1-2 — setup, onboarding & handover.

What each test is defending:

  * a school code is minted at creation, random and unguessable (S-57) —
    DPS2024 plus a browsable class list is the whole school's roster;
  * the DOB parser accepts what Indian registers actually contain and NEVER
    guesses — an unreadable value is reported, the student is still created,
    and a US-ordered date fails rather than silently swapping day and month
    (it becomes somebody's password, D-13);
  * every template maps 100% against its own importer — a template that
    drifted would send the school through the gap screen it exists to skip;
  * the readiness report names consequences ("N parents cannot log in") and a
    partial (term-wise) plan counts as ready — the mid-year guarantee;
  * work categories keep stable keys on rename and retire instead of delete —
    a school renaming "Event work" must not orphan last term's rows (D-19);
  * the class-teacher field is settable and the year-wide class-subjects read
    exists for the by-teacher lens (D-03 / D-28).
"""

import io
import uuid
from datetime import date, datetime

from openpyxl import load_workbook
from sqlalchemy import select

from app.models import User
from app.services.roster_import import parse_dob
from tests.conftest import AdminSession


def _register(client, cleanup, org_name="Setup Org"):
    email = f"op-{uuid.uuid4().hex[:12]}@example.com"
    reg = client.post("/api/v1/auth/register-org",
                      json={"org_name": org_name, "name": "Operator", "email": email,
                            "password": "supersecret1", "timezone": "Asia/Kolkata"}).json()
    cleanup["orgs"].append(uuid.UUID(reg["org"]["id"]))
    cleanup["users"].append(uuid.UUID(reg["user"]["id"]))
    return reg


def _make_super(user_id: str) -> None:
    db = AdminSession()
    try:
        db.get(User, uuid.UUID(user_id)).is_super_admin = True
        db.commit()
    finally:
        db.close()


def _headers(reg):
    return {"Authorization": f"Bearer {reg['access_token']}"}


# ── school code (D-13, S-57) ─────────────────────────────────────────────────
def test_created_school_gets_an_unguessable_code(client, cleanup):
    reg = _register(client, cleanup)
    _make_super(reg["user"]["id"])
    h = _headers(reg)
    codes = set()
    for i in range(2):
        email = f"owner-{uuid.uuid4().hex[:12]}@example.com"
        res = client.post("/api/v1/platform/orgs", headers=h, json={
            "org_name": f"Code School {i}", "admin_name": "Owner",
            "admin_email": email, "admin_password": "handover123",
            "address": "12 MG Road, Hyderabad", "state": "Telangana",
            "board": "CBSE"}).json()
        cleanup["orgs"].append(uuid.UUID(res["org"]["id"]))
        db = AdminSession()
        try:
            cleanup["users"].append(db.scalar(select(User.id).where(User.email == email)))
        finally:
            db.close()
        code = res["school_code"]
        assert code and len(code) in (6, 7, 8)
        assert not any(c in code for c in "0O1IL"), "no lookalike characters"
        codes.add(code)
        assert res["org"]["school_code"] == code
    assert len(codes) == 2, "codes are unique"


# ── DOB parsing (D-13 / Q-24) ────────────────────────────────────────────────
def test_dob_parser_reads_indian_sheets_and_never_guesses():
    today = date(2026, 8, 1)
    assert parse_dob("14/06/2014", today) == date(2014, 6, 14)
    assert parse_dob("14-06-14", today) == date(2014, 6, 14)
    assert parse_dob("14.06.2014", today) == date(2014, 6, 14)
    # A real Excel date cell arrives stringified.
    assert parse_dob("2014-06-14 00:00:00", today) == date(2014, 6, 14)
    # An Excel serial for 2014-06-14.
    assert parse_dob("41804", today) == date(2014, 6, 14)
    # Day-first ALWAYS: 06/14/2014 (US order) has month 14 → unresolved,
    # never a silent day/month swap.
    assert parse_dob("06/14/2014", today) is None
    assert parse_dob("garbage", today) is None
    assert parse_dob("14/06/2027", today) is None, "the future is not a birthday"
    assert parse_dob("14/06/1980", today) is None, "age 46 is not a student"
    assert parse_dob("14/06/1980", today, min_age=16, max_age=80) == date(1980, 6, 14), \
        "…but it is plausible staff"


def test_roster_import_stores_dob_and_reports_unreadable_ones(client, cleanup):
    reg = _register(client, cleanup)
    h = _headers(reg)
    mapping = {"full_name": "Name", "admission_no": "Adm", "date_of_birth": "DOB"}
    rows = [
        {"Name": "Asha Rao", "Adm": "A1", "DOB": "14/06/2014"},
        {"Name": "Bilal Khan", "Adm": "A2", "DOB": "not-a-date"},
        {"Name": "Chitra Nair", "Adm": "A3", "DOB": None},
    ]
    res = client.post("/api/v1/students/import/commit", headers=h,
                      json={"mapping": mapping, "rows": rows,
                            "academic_year_id": None}).json()
    assert res["created"] == 3, "an unreadable DOB never blocks the student"
    assert len(res["unresolved"]) == 1
    assert res["unresolved"][0]["student"] == "Bilal Khan"
    assert res["unresolved"][0]["value"] == "not-a-date"

    students = client.get("/api/v1/students", headers=h).json()
    by_name = {s["full_name"]: s for s in students}
    assert by_name["Asha Rao"]["date_of_birth"] == "2014-06-14"
    assert by_name["Bilal Khan"]["date_of_birth"] is None


# ── templates round-trip against their own importers (§6 ②) ──────────────────
def test_templates_map_cleanly_against_their_own_importers(client, cleanup):
    from app.services import roster_import, staff_import, syllabus_import

    reg = _register(client, cleanup)
    h = _headers(reg)

    r = client.get("/api/v1/students/import/template", headers=h)
    assert r.status_code == 200
    analysis = roster_import.analyze(r.content)
    for field in roster_import.TARGET_FIELDS:
        if field == "phone":  # covered by father/mother phone columns
            continue
        assert field in analysis["mapping"], f"template misses its own field: {field}"
    assert analysis["missing_required"] == []
    assert analysis["questions"] == []

    r = client.get("/api/v1/org/members/import/template", headers=h)
    assert r.status_code == 200
    analysis = staff_import.analyze(r.content)
    for spec in staff_import.SPECS:
        assert spec.name in analysis["mapping"]
    assert analysis["missing_required"] == []

    r = client.get("/api/v1/planner/syllabus/import/template", headers=h)
    assert r.status_code == 200
    analysis = syllabus_import.analyze_file(r.content, "syllabus-template.xlsx")
    assert analysis["mode"] == "grid"
    # The example rows demonstrate the conventions: blank chapter continues,
    # blank periods = NOT SIZED (never 1).
    units = analysis["units"]
    assert [u["title"] for u in units] == ["Number Systems", "Polynomials"]
    assert len(units[0]["topics"]) == 2, "blank chapter cell continued the chapter"
    assert units[1]["topics"][0]["est_periods"] is None, "blank periods = unsized"

    # The workbook opens and carries a note on every header.
    wb = load_workbook(io.BytesIO(client.get(
        "/api/v1/students/import/template", headers=h).content))
    ws = wb.active
    assert all(c.comment is not None for c in ws[1] if c.value)


# ── settings: D-01 mode, thresholds, work categories (D-19/S-69) ─────────────
def test_settings_round_trip_and_work_category_rules(client, cleanup):
    reg = _register(client, cleanup)
    h = _headers(reg)

    s = client.get("/api/v1/org/settings", headers=h).json()
    assert s["attendance_mode"] == "every_period"
    assert s["min_attendance_pct"] == 75
    assert s["homework_gap_days"] == 3
    keys = {c["key"] for c in s["work_categories"]}
    assert "other" in keys and "notebook_checking" in keys

    upd = client.patch("/api/v1/org/settings", headers=h, json={
        "attendance_mode": "twice_daily", "min_attendance_pct": 80,
        "homework_gap_days": 4, "state": "Telangana", "board": "CBSE",
        "work_categories": [
            # rename keeps the key (rule 1)
            {"key": "event_work", "label": "Programmes", "active": True},
            # a new school-specific category mints a stable key
            {"label": "Assembly duty", "active": True},
            {"key": "other", "label": "Other", "active": True},
        ]}).json()
    assert upd["attendance_mode"] == "twice_daily"
    assert upd["min_attendance_pct"] == 80
    cats = {c["key"]: c for c in upd["work_categories"]}
    assert cats["event_work"]["label"] == "Programmes", "same key, new word"
    assert "assembly_duty" in cats and cats["assembly_duty"]["active"]
    # Rule 2: everything the payload dropped is retired, never deleted.
    assert "notebook_checking" in cats
    assert cats["notebook_checking"]["active"] is False
    assert cats["other"]["active"] is True

    # The picker serves only active categories, renamed.
    picker = client.get("/api/v1/staff/work-types", headers=h).json()
    labels = {p["key"]: p["label"] for p in picker}
    assert labels["event_work"] == "Programmes"
    assert "notebook_checking" not in labels

    # Rule 4: `other` cannot be retired.
    upd2 = client.patch("/api/v1/org/settings", headers=h, json={
        "work_categories": [{"key": "other", "label": "Other", "active": False}]}).json()
    other = next(c for c in upd2["work_categories"] if c["key"] == "other")
    assert other["active"] is True


# ── class teacher (D-03) + the year-wide lens read (D-28/S-77) ───────────────
def test_class_teacher_assignment_and_batched_class_subjects(client, cleanup):
    reg = _register(client, cleanup)
    h = _headers(reg)
    year = client.post("/api/v1/academics/years", headers=h,
                       json={"label": "2026-27", "start_date": "2026-04-01",
                             "end_date": "2027-03-31"}).json()
    klass = client.post("/api/v1/academics/classes", headers=h,
                        json={"academic_year_id": year["id"], "name": "6",
                              "section": "A"}).json()
    subject = client.post("/api/v1/academics/subjects", headers=h,
                          json={"name": "Science"}).json()
    bulk = client.post("/api/v1/org/members/bulk", headers=h, json={"members": [
        {"username": f"t{uuid.uuid4().hex[:8]}", "password": "supersecret1",
         "role": "teacher"}]}).json()
    cleanup["users"].append(uuid.UUID(bulk["results"][0]["user_id"]))
    members = client.get("/api/v1/org/members", headers=h).json()["members"]
    teacher_mid = next(m["member_id"] for m in members
                       if m["user_id"] == bulk["results"][0]["user_id"])
    client.post("/api/v1/academics/class-subjects", headers=h,
                json={"class_id": klass["id"], "subject_id": subject["id"],
                      "teacher_member_id": teacher_mid, "periods_per_week": 5})

    updated = client.patch(f"/api/v1/academics/classes/{klass['id']}", headers=h,
                           json={"class_teacher_member_id": teacher_mid}).json()
    assert updated["class_teacher_member_id"] == teacher_mid

    rows = client.get(f"/api/v1/academics/class-subjects?year_id={year['id']}",
                      headers=h).json()
    assert len(rows) == 1
    assert rows[0]["class_label"] == "6-A"
    assert rows[0]["subject_name"] == "Science"
    assert rows[0]["teacher_member_id"] == teacher_mid


# ── the readiness report (§6 ⑤) ──────────────────────────────────────────────
def test_readiness_report_names_consequences_and_handover_sticks(client, cleanup):
    reg = _register(client, cleanup)
    _make_super(reg["user"]["id"])
    h = _headers(reg)
    email = f"owner-{uuid.uuid4().hex[:12]}@example.com"
    school = client.post("/api/v1/platform/orgs", headers=h, json={
        "org_name": "Ready School", "admin_name": "Owner",
        "admin_email": email, "admin_password": "handover123"}).json()
    org_id = school["org"]["id"]
    cleanup["orgs"].append(uuid.UUID(org_id))
    db = AdminSession()
    try:
        cleanup["users"].append(db.scalar(select(User.id).where(User.email == email)))
    finally:
        db.close()

    # Operator enters and builds a partial school: a year, a class, a student
    # with no DOB, no class teacher.
    entered = client.post(f"/api/v1/platform/orgs/{org_id}/enter", headers=h).json()
    sh = {"Authorization": f"Bearer {entered['access_token']}"}
    year = client.post("/api/v1/academics/years", headers=sh,
                       json={"label": "2026-27", "start_date": "2026-04-01",
                             "end_date": "2027-03-31"}).json()
    client.post("/api/v1/academics/classes", headers=sh,
                json={"academic_year_id": year["id"], "name": "6", "section": "A"})
    client.post("/api/v1/students", headers=sh,
                json={"admission_no": "R1", "full_name": "Kabir Shah"})

    report = client.get(f"/api/v1/platform/orgs/{org_id}/readiness", headers=h).json()
    assert report["total"] == 11
    assert report["school_code"] == school["school_code"]
    checks = {c["key"]: c for c in report["checks"]}
    assert checks["dob"]["status"] == "warn"
    assert "cannot log in" in checks["dob"]["summary"], "consequence, not field"
    assert "Kabir Shah" in checks["dob"]["items"]
    assert checks["class_teachers"]["status"] == "warn"
    assert "no owner" in checks["class_teachers"]["summary"]
    assert checks["portal"]["status"] == "ok", "code minted + portal on by default"
    assert report["handed_over_at"] is None

    done = client.post(f"/api/v1/platform/orgs/{org_id}/handover", headers=h).json()
    assert done["handed_over_at"] is not None
    again = client.post(f"/api/v1/platform/orgs/{org_id}/handover", headers=h).json()
    # Compare instants, not strings — the write comes back Z, the read +05:30.
    assert (datetime.fromisoformat(again["handed_over_at"])
            == datetime.fromisoformat(done["handed_over_at"])), "idempotent"

    # A plain admin of another org can never read a readiness report.
    other = _register(client, cleanup, org_name="Nosy Org")
    r = client.get(f"/api/v1/platform/orgs/{org_id}/readiness", headers=_headers(other))
    assert r.status_code == 403
