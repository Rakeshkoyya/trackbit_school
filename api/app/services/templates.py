"""Blank .xlsx templates for the three importers (V1-2, plan §6 ②).

Each template is generated FROM the importer's own field list, so the headers
the school fills in are — by construction — headers the importer's heuristic
maps back with zero gap questions. A template that drifted from its importer
would be worse than none: the school would fill it perfectly and still hit the
mapping screen. `test_setup_onboarding.py` asserts the round trip.

Every column carries a note (required? format?) and three example rows show the
conventions — merged-cell chapter continuation, blank periods = not sized yet,
day-first dates.
"""

import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


class TemplateColumn:
    def __init__(self, header: str, note: str, examples: tuple[str, str, str],
                 required: bool = False):
        self.header = header
        self.note = note
        self.examples = examples
        self.required = required


def build_template(sheet_name: str, columns: list[TemplateColumn]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = sheet_name
    header_font = Font(bold=True)
    required_fill = PatternFill("solid", fgColor="FFF3D6")

    for i, col in enumerate(columns, start=1):
        cell = ws.cell(row=1, column=i, value=col.header)
        cell.font = header_font
        if col.required:
            cell.fill = required_fill
        note = ("Required. " if col.required else "Optional. ") + col.note
        cell.comment = Comment(note, "TrackBit")
        width = max(len(col.header), *(len(e) for e in col.examples), 12) + 2
        ws.column_dimensions[get_column_letter(i)].width = min(width, 42)
        for r, example in enumerate(col.examples, start=2):
            ws.cell(row=r, column=i, value=example)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def roster_template() -> bytes:
    """Students. Headers chosen to hit roster_import.FIELD_HINTS exactly."""
    return build_template("Students", [
        TemplateColumn("Student Name", "The child's full name.",
                       ("Aarav Sharma", "Diya Patel", "Rohan Verma"), required=True),
        TemplateColumn("Admission No", "The school's admission / scholar number — "
                       "must be unique.", ("1021", "1022", "1023"), required=True),
        TemplateColumn("Roll No", "Roll number within the class.", ("12", "13", "14")),
        TemplateColumn("Class", "Class name exactly as created in TrackBit, e.g. 6.",
                       ("6", "6", "7")),
        TemplateColumn("Section", "Section letter, if the class has sections.",
                       ("A", "A", "B")),
        TemplateColumn("Date of Birth", "Day first: dd/mm/yyyy. This becomes the "
                       "parent's login password, so please fill it for every child.",
                       ("14/06/2014", "02/11/2013", "23/01/2014")),
        TemplateColumn("Category", "Your own grouping, e.g. Hosteller / Day scholar.",
                       ("Day scholar", "Hosteller", "Day scholar")),
        TemplateColumn("Father Name", "", ("Rajesh Sharma", "Amit Patel", "Suresh Verma")),
        TemplateColumn("Father Phone", "10-digit mobile — absence alerts and homework "
                       "notes go here.", ("9876543210", "9876543212", "9876543214")),
        TemplateColumn("Mother Name", "", ("Priya Sharma", "Meena Patel", "Lakshmi Verma")),
        TemplateColumn("Mother Phone", "", ("9876543211", "9876543213", "9876543215")),
    ])


def staff_template() -> bytes:
    """Teachers. Headers chosen to hit staff_import.SPECS hints exactly."""
    return build_template("Staff", [
        TemplateColumn("Teacher Name", "Full name.",
                       ("Anita Desai", "Vikram Rao", "Sunita Iyer"), required=True),
        TemplateColumn("Username", "Login username or employee id. Left blank, one is "
                       "generated from the name.", ("anita.desai", "vikram.rao", "")),
        TemplateColumn("Email", "", ("anita@school.example", "vikram@school.example", "")),
        TemplateColumn("Phone", "10-digit mobile.", ("9876501234", "9876501235", "")),
        TemplateColumn("Date of Birth", "Day first: dd/mm/yyyy — feeds the staff "
                       "birthday feed.", ("03/07/1988", "21/12/1990", "")),
        TemplateColumn("Subjects Taught", "One entry per class-subject, separated by "
                       "semicolons: 6-A Mathematics x6 means class 6 section A, 6 "
                       "periods a week.",
                       ("6-A Mathematics x6; 7-A Mathematics x6",
                        "6-A Science x5; 6-B Science x5",
                        "8-A English x6")),
    ])


def syllabus_template() -> bytes:
    """One sheet per class-subject. Headers hit syllabus_import.SPECS exactly.
    The examples demonstrate the two conventions the importer honours: a blank
    chapter cell continues the previous chapter, and a blank periods cell means
    NOT SIZED YET (the mid-year case) — never 1."""
    return build_template("Syllabus", [
        TemplateColumn("Chapter", "Chapter / unit name. Leave blank to continue the "
                       "chapter above (like merged cells).",
                       ("Number Systems", "", "Polynomials"), required=True),
        TemplateColumn("Topic", "One row per topic.",
                       ("Rational numbers", "Irrational numbers",
                        "Introduction to polynomials"), required=True),
        TemplateColumn("Periods", "How many periods this topic needs. Leave blank if "
                       "you haven't sized it yet — blank means 'not sized', never 1.",
                       ("3", "2", "")),
        TemplateColumn("Term", "Which term this chapter belongs to, exactly as named "
                       "in TrackBit (e.g. Term 1). Blank continues the term above.",
                       ("Term 1", "", "Term 2")),
    ])
