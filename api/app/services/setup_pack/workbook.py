"""Generate the blank setup pack the school fills in (SETUP-REDESIGN-PLAN §4).

One workbook, one download, ten sheets plus a Read Me — replacing the twenty-three
separate files a mid-sized school needs today (`test_doc/new_org/` writes twenty
syllabus files for a four-class school).

Two deliberate choices:

  * **The blank pack ships with headers only, no example rows.** `templates.py`
    seeds three example rows per template, which is friendly right up until a school
    returns the file with "Aarav Sharma" still in it and we import him. Guidance
    lives in the cell comments and the Read Me instead, where it cannot be mistaken
    for data. `build_pack(sample=True)` produces the filled version for tests and
    for `test_doc/new_org`.
  * **Fixed value sets become real dropdowns.** A `Role` column that can only say
    Teacher or Admin is worth more as a dropdown than as a paragraph nobody reads.
"""

import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.worksheet import Worksheet

from app.services.setup_pack.specs import READ_ME, READ_ME_NOTES, SHEETS, SheetSpec

# How far down a sheet the dropdowns reach. A school with more than this many
# students simply types past the styling — the parser does not care.
VALIDATION_ROWS = 2000

_HEADER_FONT = Font(bold=True, color="FFFFFF")
_HEADER_FILL = PatternFill("solid", fgColor="1F3A5F")
_REQUIRED_FILL = PatternFill("solid", fgColor="FFF3D6")
_LABEL_FONT = Font(bold=True)
_TITLE_FONT = Font(bold=True, size=14)
_MUTED_FONT = Font(color="666666")


def _comment(note: str, example: str, required: bool) -> Comment:
    body = "REQUIRED. " if required else "Optional. "
    body += note or ""
    if example:
        body += f"\n\nExample: {example}"
    c = Comment(body.strip(), "TrackBit")
    c.width, c.height = 260, 130
    return c


def _dropdown(ws: Worksheet, choices: tuple[str, ...], cell_range: str) -> None:
    """Excel caps an inline list at 255 characters — beyond that the file opens
    with a repair prompt, which reads to the school as 'TrackBit sent a broken
    file'. Skip the dropdown rather than risk it; the comment still says what is
    allowed."""
    joined = ",".join(choices)
    if not choices or len(joined) > 240:
        return
    dv = DataValidation(type="list", formula1=f'"{joined}"', allow_blank=True)
    dv.error = "Pick one of the listed values."
    dv.promptTitle = "Choose a value"
    ws.add_data_validation(dv)
    dv.add(cell_range)


def _width(*values: str) -> int:
    return min(max((len(v) for v in values if v), default=12) + 4, 44)


def _write_table(ws: Worksheet, spec: SheetSpec, sample: bool) -> None:
    """Header row + optional sample rows. Required headers are tinted so the
    school can see at a glance what it cannot leave out."""
    for i, col in enumerate(spec.columns, start=1):
        cell = ws.cell(row=1, column=i, value=col.header)
        cell.font = _LABEL_FONT if col.required else _HEADER_FONT
        cell.fill = _REQUIRED_FILL if col.required else _HEADER_FILL
        cell.alignment = Alignment(vertical="center")
        cell.comment = _comment(col.note, col.example, col.required)
        letter = get_column_letter(i)
        ws.column_dimensions[letter].width = _width(col.header, col.example)
        _dropdown(ws, col.choices, f"{letter}2:{letter}{VALIDATION_ROWS}")

    if sample:
        for r, row in enumerate(spec.samples, start=2):
            for c, value in enumerate(row, start=1):
                ws.cell(row=r, column=c, value=value)
    ws.freeze_panes = "A2"


def _write_key_value(ws: Worksheet, spec: SheetSpec, sample: bool,
                     overrides: dict[str, str]) -> None:
    """The School sheet: Field | Value down the page. Fifteen one-off facts read
    better as rows than as one very wide table row."""
    for i, (header, width) in enumerate(((("Field"), 30), ("Value", 40)), start=1):
        cell = ws.cell(row=1, column=i, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        ws.column_dimensions[get_column_letter(i)].width = width

    for r, setting in enumerate(spec.settings, start=2):
        label = ws.cell(row=r, column=1, value=setting.label)
        label.font = _LABEL_FONT
        if setting.required:
            label.fill = _REQUIRED_FILL
        label.comment = _comment(setting.note, setting.example, setting.required)
        value = overrides.get(setting.key)
        if value is None and sample:
            value = setting.example
        if value is not None:
            ws.cell(row=r, column=2, value=value)
        _dropdown(ws, setting.choices, f"B{r}:B{r}")
    ws.freeze_panes = "A2"


def _write_read_me(ws: Worksheet, school_name: str | None) -> None:
    ws.column_dimensions["A"].width = 4
    ws.column_dimensions["B"].width = 104
    row = 1

    def line(text: str = "", font: Font | None = None) -> None:
        nonlocal row
        cell = ws.cell(row=row, column=2, value=text)
        if font:
            cell.font = font
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        row += 1

    line(f"TrackBit setup pack — {school_name}" if school_name
         else "TrackBit setup pack", _TITLE_FONT)
    line("Everything TrackBit needs to run this school. Fill it in, send it back, "
         "and the school is live.", _MUTED_FONT)
    line()

    line("THE SHEETS", _LABEL_FONT)
    for spec in SHEETS:
        marker = "required" if spec.required else "optional"
        line(f"• {spec.title}  ({marker}) — {spec.blurb}")
    line()

    for note in READ_ME_NOTES:
        line(note.heading.upper(), _LABEL_FONT)
        line(note.body)
        for item in note.lines:
            line(f"   • {item}")
        line()

    line("A tinted column header means the column is required. Hover any header "
         "for a note on what goes in it.", _MUTED_FONT)
    line("Questions on any of this — ask us before guessing. A blank is always "
         "safer than a guess.", _MUTED_FONT)


def build_pack(school_name: str | None = None, *, sample: bool = False,
               overrides: dict[str, str] | None = None) -> bytes:
    """The workbook, as bytes.

    `sample=True` fills every sheet with the spec's example rows — that is the
    fixture the round-trip test parses back, and the shape `test_doc/new_org`
    will emit. The pack a school receives is always blank.

    `overrides` pre-fills named School settings (the operator already typed the
    school's name and state when creating the org, so the school should not have
    to type them again).
    """
    wb = Workbook()
    wb.remove(wb.active)
    _write_read_me(wb.create_sheet(READ_ME), school_name)
    for spec in SHEETS:
        ws = wb.create_sheet(spec.title)
        if spec.is_key_value:
            _write_key_value(ws, spec, sample, overrides or {})
        else:
            _write_table(ws, spec, sample)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def filename_for(school_name: str | None) -> str:
    """A filename the operator can find again in a folder of thirty of them."""
    safe = "".join(ch for ch in (school_name or "")
                   if ch.isalnum() or ch in " -_").strip()
    slug = safe.replace(" ", "-") if safe else "School"
    return f"TrackBit-Setup-Pack-{slug}.xlsx"
